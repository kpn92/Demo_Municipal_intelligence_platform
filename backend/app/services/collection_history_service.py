from datetime import date
from typing import Optional
from sqlalchemy.orm import Session
from sqlalchemy import text
import io

from app.schemas.collection_history import CollectionEventRead, HistoryStatsRead, DayCount, TypeCount, BinNotCollected


def get_collection_events(
    db: Session,
    date_from: date,
    date_to: date,
    bin_type: Optional[str] = None,
    area_name: Optional[str] = None,
    vehicle_id: Optional[int] = None,
) -> list[CollectionEventRead]:
    filters = [
        "e.event_type = 'collection'",
        "e.timestamp::date BETWEEN :date_from AND :date_to",
    ]
    params: dict = {"date_from": date_from, "date_to": date_to}

    if bin_type:
        filters.append("b.type = :bin_type")
        params["bin_type"] = bin_type
    if area_name:
        filters.append("b.area_name = :area_name")
        params["area_name"] = area_name
    if vehicle_id:
        filters.append("e.vehicle_id = :vehicle_id")
        params["vehicle_id"] = vehicle_id

    where = " AND ".join(filters)
    rows = db.execute(
        text(f"""
            SELECT
                e.id,
                e.waste_bin_id        AS bin_id,
                b.bin_code,
                b.type                AS bin_type,
                e.vehicle_id,
                v.plate_number        AS vehicle_plate,
                v.vehicle_code        AS vehicle_name,
                e.timestamp           AS collected_at,
                e.fill_level_before,
                b.area_name,
                b.area_code,
                ST_Y(b.location::geometry) AS lat,
                ST_X(b.location::geometry) AS lng
            FROM collection_execution_event e
            JOIN waste_bin  b ON b.id = e.waste_bin_id
            LEFT JOIN vehicle v ON v.id = e.vehicle_id
            WHERE {where}
            ORDER BY e.timestamp DESC
            LIMIT 2000
        """),
        params,
    ).mappings()
    return [CollectionEventRead(**dict(r)) for r in rows]


def get_history_stats(
    db: Session,
    date_from: date,
    date_to: date,
    bin_type: Optional[str] = None,
    area_name: Optional[str] = None,
) -> HistoryStatsRead:
    filters = [
        "e.event_type = 'collection'",
        "e.timestamp::date BETWEEN :date_from AND :date_to",
    ]
    params: dict = {"date_from": date_from, "date_to": date_to}

    if bin_type:
        filters.append("b.type = :bin_type")
        params["bin_type"] = bin_type
    if area_name:
        filters.append("b.area_name = :area_name")
        params["area_name"] = area_name

    where = " AND ".join(filters)

    # Aggregate stats
    agg = db.execute(
        text(f"""
            SELECT
                COUNT(*)                          AS total_collections,
                COALESCE(AVG(e.fill_level_before), 0) AS avg_fill_level
            FROM collection_execution_event e
            JOIN waste_bin b ON b.id = e.waste_bin_id
            WHERE {where}
        """),
        params,
    ).mappings().fetchone()

    # Most active vehicle
    top_vehicle_row = db.execute(
        text(f"""
            SELECT v.plate_number, COUNT(*) AS cnt
            FROM collection_execution_event e
            JOIN waste_bin b ON b.id = e.waste_bin_id
            LEFT JOIN vehicle v ON v.id = e.vehicle_id
            WHERE {where} AND v.plate_number IS NOT NULL
            GROUP BY v.plate_number
            ORDER BY cnt DESC
            LIMIT 1
        """),
        params,
    ).mappings().fetchone()

    # Bins not collected in period
    bin_filter = ""
    if bin_type:
        bin_filter += " AND type = :bin_type"
    if area_name:
        bin_filter += " AND area_name = :area_name"

    not_collected_row = db.execute(
        text(f"""
            SELECT COUNT(*) AS cnt
            FROM waste_bin
            WHERE is_active IS TRUE {bin_filter}
              AND id NOT IN (
                  SELECT DISTINCT e.waste_bin_id
                  FROM collection_execution_event e
                  JOIN waste_bin b ON b.id = e.waste_bin_id
                  WHERE {where}
              )
        """),
        params,
    ).mappings().fetchone()

    # Collections by day
    by_day_rows = db.execute(
        text(f"""
            SELECT e.timestamp::date AS day, COUNT(*) AS cnt
            FROM collection_execution_event e
            JOIN waste_bin b ON b.id = e.waste_bin_id
            WHERE {where}
            GROUP BY day
            ORDER BY day
        """),
        params,
    ).mappings().all()

    # Collections by type
    by_type_rows = db.execute(
        text(f"""
            SELECT b.type AS bin_type, COUNT(*) AS cnt
            FROM collection_execution_event e
            JOIN waste_bin b ON b.id = e.waste_bin_id
            WHERE {where}
            GROUP BY b.type
        """),
        params,
    ).mappings().all()

    return HistoryStatsRead(
        total_collections=agg["total_collections"] if agg else 0,
        avg_fill_level=round(float(agg["avg_fill_level"]) if agg else 0, 1),
        most_active_vehicle=top_vehicle_row["plate_number"] if top_vehicle_row else None,
        bins_not_collected=not_collected_row["cnt"] if not_collected_row else 0,
        collections_by_day=[DayCount(date=str(r["day"]), count=r["cnt"]) for r in by_day_rows],
        collections_by_type=[TypeCount(bin_type=r["bin_type"], count=r["cnt"]) for r in by_type_rows],
    )


def get_bins_not_collected(
    db: Session,
    date_from: date,
    date_to: date,
    bin_type: Optional[str] = None,
    area_name: Optional[str] = None,
) -> list[BinNotCollected]:
    bin_filters = ["b.is_active IS TRUE"]
    params: dict = {"date_from": date_from, "date_to": date_to}

    if bin_type:
        bin_filters.append("b.type = :bin_type")
        params["bin_type"] = bin_type
    if area_name:
        bin_filters.append("b.area_name = :area_name")
        params["area_name"] = area_name

    bin_where = " AND ".join(bin_filters)

    rows = db.execute(
        text(f"""
            SELECT
                b.id                             AS bin_id,
                b.bin_code,
                b.type                           AS bin_type,
                b.area_name,
                ST_Y(b.location::geometry)       AS lat,
                ST_X(b.location::geometry)       AS lng,
                last_ev.last_collected_at,
                CASE
                    WHEN last_ev.last_collected_at IS NULL THEN NULL
                    ELSE (CURRENT_DATE - last_ev.last_collected_at::date)
                END                              AS days_since_last
            FROM waste_bin b
            LEFT JOIN (
                SELECT waste_bin_id, MAX(timestamp) AS last_collected_at
                FROM collection_execution_event
                WHERE event_type = 'collection'
                GROUP BY waste_bin_id
            ) last_ev ON last_ev.waste_bin_id = b.id
            WHERE {bin_where}
              AND b.id NOT IN (
                  SELECT DISTINCT e.waste_bin_id
                  FROM collection_execution_event e
                  WHERE e.event_type = 'collection'
                    AND e.timestamp::date BETWEEN :date_from AND :date_to
              )
            ORDER BY days_since_last DESC NULLS LAST
        """),
        params,
    ).mappings().all()

    return [BinNotCollected(**dict(r)) for r in rows]


def get_bin_stats_geojson(
    db: Session,
    date_from: date,
    date_to: date,
    bin_type: Optional[str] = None,
    area_name: Optional[str] = None,
) -> dict:
    filters = [
        "e.event_type = 'collection'",
        "e.timestamp::date BETWEEN :date_from AND :date_to",
    ]
    params: dict = {"date_from": date_from, "date_to": date_to}

    if bin_type:
        filters.append("b.type = :bin_type")
        params["bin_type"] = bin_type
    if area_name:
        filters.append("b.area_name = :area_name")
        params["area_name"] = area_name

    where = " AND ".join(filters)

    rows = db.execute(
        text(f"""
            SELECT
                b.id                             AS bin_id,
                b.bin_code,
                b.type                           AS bin_type,
                b.area_name,
                b.area_code,
                COUNT(e.id)                      AS collection_count,
                MAX(e.timestamp)                 AS last_collected_at,
                COALESCE(AVG(e.fill_level_before), 0) AS avg_fill_level,
                ST_Y(b.location::geometry)       AS lat,
                ST_X(b.location::geometry)       AS lng
            FROM collection_execution_event e
            JOIN waste_bin b ON b.id = e.waste_bin_id
            WHERE {where}
            GROUP BY b.id, b.bin_code, b.type, b.area_name, b.area_code,
                     b.location
        """),
        params,
    ).mappings().all()

    features = []
    for r in rows:
        features.append({
            "type": "Feature",
            "geometry": {"type": "Point", "coordinates": [float(r["lng"]), float(r["lat"])]},
            "properties": {
                "bin_id": r["bin_id"],
                "bin_code": r["bin_code"],
                "bin_type": r["bin_type"],
                "area_name": r["area_name"],
                "area_code": r["area_code"],
                "collection_count": r["collection_count"],
                "last_collected_at": str(r["last_collected_at"]) if r["last_collected_at"] else None,
                "avg_fill_level": round(float(r["avg_fill_level"]), 1),
            },
        })

    return {"type": "FeatureCollection", "features": features}


def export_history_excel(
    db: Session,
    date_from: date,
    date_to: date,
    bin_type: Optional[str] = None,
    area_name: Optional[str] = None,
) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    events = get_collection_events(db, date_from, date_to, bin_type, area_name)
    not_collected = get_bins_not_collected(db, date_from, date_to, bin_type, area_name)

    wb = openpyxl.Workbook()

    # Sheet 1: Detailed events
    ws1 = wb.active
    ws1.title = "Αναλυτικά"
    header_fill = PatternFill("solid", fgColor="1a2744")
    header_font = Font(bold=True, color="FFFFFF")
    headers1 = ["ID", "Κωδ. Κάδου", "Τύπος", "Περιοχή", "Κωδ. Περιοχής",
                 "Πινακίδα", "Ημ/νία Αποκομιδής", "Πλήρωση %"]
    for col, h in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row_idx, ev in enumerate(events, 2):
        ws1.cell(row=row_idx, column=1, value=ev.id)
        ws1.cell(row=row_idx, column=2, value=ev.bin_code)
        ws1.cell(row=row_idx, column=3, value=ev.bin_type)
        ws1.cell(row=row_idx, column=4, value=ev.area_name)
        ws1.cell(row=row_idx, column=5, value=ev.area_code)
        ws1.cell(row=row_idx, column=6, value=ev.vehicle_plate)
        ws1.cell(row=row_idx, column=7, value=str(ev.collected_at)[:16])
        ws1.cell(row=row_idx, column=8, value=ev.fill_level_before)
    for col in ws1.columns:
        ws1.column_dimensions[col[0].column_letter].width = 18

    # Sheet 2: Area summary
    ws2 = wb.create_sheet("Ανά Περιοχή")
    area_agg: dict = {}
    for ev in events:
        key = ev.area_name or "-"
        if key not in area_agg:
            area_agg[key] = {"count": 0, "fill_sum": 0, "fill_n": 0}
        area_agg[key]["count"] += 1
        if ev.fill_level_before is not None:
            area_agg[key]["fill_sum"] += ev.fill_level_before
            area_agg[key]["fill_n"] += 1
    headers2 = ["Περιοχή", "Αποκομιδές", "Μ.Ο. Πλήρωσης %"]
    for col, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row_idx, (area, data) in enumerate(sorted(area_agg.items()), 2):
        avg_fill = round(data["fill_sum"] / data["fill_n"], 1) if data["fill_n"] else 0
        ws2.cell(row=row_idx, column=1, value=area)
        ws2.cell(row=row_idx, column=2, value=data["count"])
        ws2.cell(row=row_idx, column=3, value=avg_fill)
    for col in ws2.columns:
        ws2.column_dimensions[col[0].column_letter].width = 22

    # Sheet 3: Bins not collected
    ws3 = wb.create_sheet("Κάδοι χωρίς Αποκομιδή")
    headers3 = ["ID", "Κωδ. Κάδου", "Τύπος", "Περιοχή", "Τελ. Αποκομιδή", "Ημέρες"]
    for col, h in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row_idx, b in enumerate(not_collected, 2):
        ws3.cell(row=row_idx, column=1, value=b.bin_id)
        ws3.cell(row=row_idx, column=2, value=b.bin_code)
        ws3.cell(row=row_idx, column=3, value=b.bin_type)
        ws3.cell(row=row_idx, column=4, value=b.area_name)
        ws3.cell(row=row_idx, column=5, value=str(b.last_collected_at)[:10] if b.last_collected_at else "Ποτέ")
        ws3.cell(row=row_idx, column=6, value=b.days_since_last)
    for col in ws3.columns:
        ws3.column_dimensions[col[0].column_letter].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _register_greek_font() -> str:
    """Register a Unicode TTF font that supports Greek. Returns the font name."""
    import os
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_name = "ArialUnicode"
    if font_name in pdfmetrics.getRegisteredFontNames():
        return font_name

    candidates = [
        os.path.join(os.path.dirname(__file__), "..", "static", "fonts", "arial.ttf"),
        os.path.join(os.path.dirname(__file__), "..", "static", "fonts", "DejaVuSans.ttf"),
        "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    ]
    for path in candidates:
        path = os.path.normpath(path)
        if os.path.isfile(path):
            pdfmetrics.registerFont(TTFont(font_name, path))
            return font_name

    return "Helvetica"


def export_history_pdf(
    db: Session,
    date_from: date,
    date_to: date,
    bin_type: Optional[str] = None,
    area_name: Optional[str] = None,
) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm

    font = _register_greek_font()
    font_bold = font  # TTFont doesn't have separate bold; use same font for now

    stats = get_history_stats(db, date_from, date_to, bin_type, area_name)
    not_collected = get_bins_not_collected(db, date_from, date_to, bin_type, area_name)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    navy = colors.HexColor("#1a2744")

    title_style = ParagraphStyle("MIPTitle", fontName=font, fontSize=16,
                                 textColor=navy, spaceAfter=6)
    h2_style = ParagraphStyle("MIPH2", fontName=font, fontSize=12,
                              textColor=navy, spaceBefore=12, spaceAfter=4)
    normal_style = ParagraphStyle("MIPNormal", fontName=font, fontSize=10,
                                  spaceAfter=4)
    footer_style = ParagraphStyle("MIPFooter", fontName=font, fontSize=8,
                                  textColor=colors.grey)

    story = []

    story.append(Paragraph("Δήμος Κερατσινίου-Δραπετσώνας", title_style))
    story.append(Paragraph("Ιστορικό Αποκομιδής Κάδων", h2_style))
    period_label = f"Περίοδος: {date_from} - {date_to}"
    if bin_type:
        period_label += f" | Τύπος: {bin_type}"
    if area_name:
        period_label += f" | Περιοχή: {area_name}"
    story.append(Paragraph(period_label, normal_style))
    story.append(Spacer(1, 0.4*cm))

    ts_header = [
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
    ]

    # KPI table
    kpi_data = [
        ["Σύνολο Αποκομιδών", "Μ.Ο. Πλήρωσης", "Κορυφαίο Όχημα", "Κάδοι χωρίς Αποκομιδή"],
        [
            str(stats.total_collections),
            f"{stats.avg_fill_level}%",
            stats.most_active_vehicle or "-",
            str(stats.bins_not_collected),
        ],
    ]
    kpi_table = Table(kpi_data, colWidths=[4.2*cm]*4)
    kpi_table.setStyle(TableStyle(ts_header))
    story.append(kpi_table)
    story.append(Spacer(1, 0.5*cm))

    # Collections by type
    if stats.collections_by_type:
        story.append(Paragraph("Αποκομιδές ανά Τύπο Κάδου", h2_style))
        type_map = {"mixed": "Μικτά", "recycling": "Ανακύκλωση", "organic": "Οργανικά"}
        type_data = [["Τύπος", "Αποκομιδές"]]
        for tc in stats.collections_by_type:
            type_data.append([type_map.get(tc.bin_type, tc.bin_type), str(tc.count)])
        type_table = Table(type_data, colWidths=[8*cm, 8*cm])
        type_table.setStyle(TableStyle(ts_header))
        story.append(type_table)
        story.append(Spacer(1, 0.5*cm))

    # Bins not collected
    story.append(Paragraph(f"Κάδοι χωρίς Αποκομιδή ({len(not_collected)})", h2_style))
    if not_collected:
        type_map = {"mixed": "Μικτά", "recycling": "Ανακύκλωση", "organic": "Οργανικά"}
        nc_data = [["Κωδ. Κάδου", "Τύπος", "Περιοχή", "Τελ. Αποκομιδή", "Ημέρες"]]
        for b in not_collected[:50]:
            nc_data.append([
                b.bin_code,
                type_map.get(b.bin_type, b.bin_type),
                b.area_name or "-",
                str(b.last_collected_at)[:10] if b.last_collected_at else "Ποτέ",
                str(b.days_since_last) if b.days_since_last is not None else "-",
            ])
        nc_ts = list(ts_header) + [("FONTSIZE", (0, 0), (-1, -1), 8)]
        nc_table = Table(nc_data, colWidths=[3.5*cm, 3*cm, 4*cm, 4*cm, 2.5*cm])
        nc_table.setStyle(TableStyle(nc_ts))
        story.append(nc_table)
    else:
        story.append(Paragraph(
            "Όλοι οι κάδοι έχουν αποκομιστεί στην επιλεγμένη περίοδο.", normal_style
        ))

    story.append(Spacer(1, 0.8*cm))
    from datetime import datetime
    story.append(Paragraph(
        f"Δήμος Κερατσινίου-Δραπετσώνας | Εκτύπωση: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
        footer_style,
    ))

    doc.build(story)
    return buf.getvalue()
